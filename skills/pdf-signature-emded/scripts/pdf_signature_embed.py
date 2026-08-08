#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any


ANCHORS = {"top-left", "top-right", "bottom-left", "bottom-right", "center"}
ROTATIONS = {0, 90, 180, 270}


class ConfigError(Exception):
    pass


def load_fitz():
    try:
        import fitz  # type: ignore
    except ImportError as exc:
        raise SystemExit(
            "PyMuPDF is required. Install it with: python3 -m pip install pymupdf"
        ) from exc
    return fitz


def load_docx():
    try:
        import docx  # type: ignore
        from docx.shared import Inches  # type: ignore
    except ImportError as exc:
        raise SystemExit(
            "python-docx is required for Word documents. "
            "Install it with: python3 -m pip install python-docx"
        ) from exc
    return docx, Inches


def read_json(path: Path) -> dict[str, Any]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ConfigError(f"{path}: invalid JSON: {exc}") from exc
    if not isinstance(data, dict):
        raise ConfigError(f"{path}: config root must be an object")
    return data


def resolve_path(config_path: Path, value: str) -> Path:
    path = Path(value).expanduser()
    if path.is_absolute():
        return path
    return (config_path.parent / path).resolve()


def require_string(obj: dict[str, Any], key: str, context: str) -> str:
    value = obj.get(key)
    if not isinstance(value, str) or not value:
        raise ConfigError(f"{context}: {key} must be a non-empty string")
    return value


def require_number(obj: dict[str, Any], key: str, context: str) -> float:
    value = obj.get(key)
    if not isinstance(value, (int, float)):
        raise ConfigError(f"{context}: {key} must be a number")
    return float(value)


def validate_string_list(value: Any, context: str) -> None:
    if value is None:
        return
    if not isinstance(value, list) or any(not isinstance(item, str) for item in value):
        raise ConfigError(f"{context}: must be an array of strings")


def portable_path(config_path: Path, path: Path) -> str:
    resolved = path.expanduser().resolve()
    try:
        return str(resolved.relative_to(config_path.parent.resolve()))
    except ValueError:
        return str(resolved)


def safe_asset_name(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", name).strip("-")
    return cleaned or "signature"


@dataclass(frozen=True)
class DocumentSpec:
    id: str
    path: Path
    output: Path
    kind: str


@dataclass(frozen=True)
class SignatureSpec:
    id: str
    path: Path


@dataclass(frozen=True)
class PlacementSpec:
    document: str
    signature: str
    page: int | None
    anchor: str
    x: float | None
    y: float | None
    width: float
    height: float | None
    rotation: int
    label: str | None
    mode: str | None
    placeholder: str | None
    paragraph: str | None


@dataclass(frozen=True)
class SignatureConfig:
    documents: dict[str, DocumentSpec]
    signatures: dict[str, SignatureSpec]
    placements: list[PlacementSpec]


def parse_config(config_path: Path, check_files: bool = True) -> SignatureConfig:
    data = read_json(config_path)
    if data.get("version") != 1:
        raise ConfigError("version must be 1")
    validate_string_list(data.get("customization_instructions"), "customization_instructions")

    documents_raw = data.get("documents")
    signatures_raw = data.get("signatures")
    placements_raw = data.get("placements")
    if not isinstance(documents_raw, list) or not documents_raw:
        raise ConfigError("documents must be a non-empty array")
    if not isinstance(signatures_raw, list) or not signatures_raw:
        raise ConfigError("signatures must be a non-empty array")
    if not isinstance(placements_raw, list) or not placements_raw:
        raise ConfigError("placements must be a non-empty array")

    documents: dict[str, DocumentSpec] = {}
    for index, item in enumerate(documents_raw):
        context = f"documents[{index}]"
        if not isinstance(item, dict):
            raise ConfigError(f"{context}: must be an object")
        doc_id = require_string(item, "id", context)
        if doc_id in documents:
            raise ConfigError(f"{context}: duplicate document id {doc_id!r}")
        path = resolve_path(config_path, require_string(item, "path", context))
        output = resolve_path(config_path, require_string(item, "output", context))
        doc_type = item.get("type")
        if doc_type is None:
            suffix = path.suffix.lower()
            if suffix == ".pdf":
                doc_type = "pdf"
            elif suffix == ".docx":
                doc_type = "docx"
            else:
                raise ConfigError(
                    f"{context}: type is required for unsupported extension {suffix!r}"
                )
        if doc_type not in {"pdf", "docx"}:
            raise ConfigError(f"{context}: type must be 'pdf' or 'docx'")
        if check_files and not path.exists():
            raise ConfigError(f"{context}: document not found: {path}")
        documents[doc_id] = DocumentSpec(doc_id, path, output, str(doc_type))

    signatures: dict[str, SignatureSpec] = {}
    for index, item in enumerate(signatures_raw):
        context = f"signatures[{index}]"
        if not isinstance(item, dict):
            raise ConfigError(f"{context}: must be an object")
        sig_id = require_string(item, "id", context)
        if sig_id in signatures:
            raise ConfigError(f"{context}: duplicate signature id {sig_id!r}")
        path = resolve_path(config_path, require_string(item, "path", context))
        if check_files and not path.exists():
            raise ConfigError(f"{context}: signature image not found: {path}")
        validate_string_list(
            item.get("customization_instructions"),
            f"{context}.customization_instructions",
        )
        signatures[sig_id] = SignatureSpec(sig_id, path)

    placements: list[PlacementSpec] = []
    for index, item in enumerate(placements_raw):
        context = f"placements[{index}]"
        if not isinstance(item, dict):
            raise ConfigError(f"{context}: must be an object")
        document = require_string(item, "document", context)
        signature = require_string(item, "signature", context)
        if document not in documents:
            raise ConfigError(f"{context}: unknown document id {document!r}")
        if signature not in signatures:
            raise ConfigError(f"{context}: unknown signature id {signature!r}")
        doc_kind = documents[document].kind
        page = item.get("page")
        if doc_kind == "pdf":
            if not isinstance(page, int) or page < 1:
                raise ConfigError(f"{context}: page must be a 1-based integer")
        elif page is not None and (not isinstance(page, int) or page < 1):
            raise ConfigError(f"{context}: page must be a 1-based integer")
        anchor = item.get("anchor", "bottom-left")
        if anchor not in ANCHORS:
            raise ConfigError(f"{context}: anchor must be one of {sorted(ANCHORS)}")
        x = None
        y = None
        if doc_kind == "pdf":
            x = require_number(item, "x", context)
            y = require_number(item, "y", context)
        elif "x" in item or "y" in item:
            x = require_number(item, "x", context)
            y = require_number(item, "y", context)
        width = require_number(item, "width", context)
        if width <= 0:
            raise ConfigError(f"{context}: width must be positive")
        height_value = item.get("height")
        height = None
        if height_value is not None:
            if not isinstance(height_value, (int, float)) or height_value <= 0:
                raise ConfigError(f"{context}: height must be a positive number")
            height = float(height_value)
        rotation = item.get("rotation", 0)
        if rotation not in ROTATIONS:
            raise ConfigError(f"{context}: rotation must be one of {sorted(ROTATIONS)}")
        label = item.get("label")
        if label is not None and not isinstance(label, str):
            raise ConfigError(f"{context}: label must be a string")
        mode = item.get("mode")
        placeholder = item.get("placeholder")
        paragraph = item.get("paragraph")
        if doc_kind == "docx":
            if mode is None:
                if isinstance(placeholder, str):
                    mode = "placeholder"
                elif isinstance(paragraph, str):
                    mode = "after-paragraph"
                else:
                    mode = "append"
            if mode not in {"placeholder", "after-paragraph", "append"}:
                raise ConfigError(
                    f"{context}: mode must be placeholder, after-paragraph, or append"
                )
            if mode == "placeholder" and not isinstance(placeholder, str):
                raise ConfigError(f"{context}: placeholder must be a string")
            if mode == "after-paragraph" and not isinstance(paragraph, str):
                raise ConfigError(f"{context}: paragraph must be a string")
        else:
            if mode is not None:
                raise ConfigError(f"{context}: mode is only supported for docx")
            if placeholder is not None or paragraph is not None:
                raise ConfigError(
                    f"{context}: placeholder and paragraph are only supported for docx"
                )
        placements.append(
            PlacementSpec(
                document=document,
                signature=signature,
                page=page,
                anchor=anchor,
                x=x,
                y=y,
                width=width,
                height=height,
                rotation=int(rotation),
                label=label,
                mode=str(mode) if mode is not None else None,
                placeholder=placeholder,
                paragraph=paragraph,
            )
        )

    return SignatureConfig(documents, signatures, placements)


def image_size(fitz: Any, image_path: Path) -> tuple[float, float]:
    pixmap = fitz.Pixmap(str(image_path))
    try:
        return float(pixmap.width), float(pixmap.height)
    finally:
        pixmap = None


def placement_rect(
    fitz: Any,
    page_height: float,
    image_width: float,
    image_height: float,
    placement: PlacementSpec,
) -> Any:
    width = placement.width
    height = placement.height or (width * image_height / image_width)

    if placement.x is None or placement.y is None:
        raise ConfigError("PDF placements require x and y coordinates")
    x = placement.x
    y = placement.y
    if placement.anchor == "bottom-left":
        left, bottom = x, y
    elif placement.anchor == "bottom-right":
        left, bottom = x - width, y
    elif placement.anchor == "top-left":
        left, bottom = x, y - height
    elif placement.anchor == "top-right":
        left, bottom = x - width, y - height
    else:
        left, bottom = x - width / 2, y - height / 2

    top = page_height - bottom - height
    return fitz.Rect(left, top, left + width, top + height)


def validate_with_pdf_metadata(config_path: Path) -> SignatureConfig:
    config = parse_config(config_path)
    if not any(document.kind == "pdf" for document in config.documents.values()):
        return config
    fitz = load_fitz()
    opened: dict[str, Any] = {}
    try:
        for doc_id, document in config.documents.items():
            if document.kind == "pdf":
                opened[doc_id] = fitz.open(str(document.path))
        for index, placement in enumerate(config.placements):
            document = config.documents[placement.document]
            if document.kind != "pdf":
                continue
            doc = opened[placement.document]
            if placement.page is None:
                raise ConfigError(f"placements[{index}]: page is required for PDFs")
            if placement.page > doc.page_count:
                raise ConfigError(
                    f"placements[{index}]: page {placement.page} exceeds "
                    f"{placement.document!r} page count {doc.page_count}"
                )
            image_w, image_h = image_size(
                fitz, config.signatures[placement.signature].path
            )
            page = doc[placement.page - 1]
            rect = placement_rect(fitz, page.rect.height, image_w, image_h, placement)
            if rect.width <= 0 or rect.height <= 0:
                raise ConfigError(f"placements[{index}]: computed rectangle is empty")
            if not page.rect.intersects(rect):
                raise ConfigError(
                    f"placements[{index}]: computed rectangle is outside the page"
                )
    finally:
        for doc in opened.values():
            doc.close()
    return config


def command_init_config(args: argparse.Namespace) -> int:
    pdfs = [Path(value) for value in args.pdf]
    signatures = [Path(value) for value in args.signature]
    output_path = Path(args.output)
    base = output_path.parent.resolve()

    def portable(path: Path) -> str:
        resolved = path.expanduser().resolve()
        try:
            return str(resolved.relative_to(base))
        except ValueError:
            return str(resolved)

    documents = []
    for index, pdf in enumerate(pdfs, start=1):
        stem = pdf.stem or f"document-{index}"
        suffix = pdf.suffix.lower()
        doc_type = "docx" if suffix == ".docx" else "pdf"
        documents.append(
            {
                "id": stem if len(pdfs) == 1 else f"{stem}-{index}",
                "type": doc_type,
                "path": portable(pdf),
                "output": f"signed/{pdf.stem or f'document-{index}'}-signed{pdf.suffix.lower() or '.pdf'}",
            }
        )
    sig_items = []
    for index, signature in enumerate(signatures, start=1):
        stem = signature.stem or f"signature-{index}"
        sig_items.append(
            {
                "id": stem if len(signatures) == 1 else f"{stem}-{index}",
                "path": portable(signature),
            }
        )

    if documents[0]["type"] == "docx":
        placement = {
            "document": documents[0]["id"],
            "signature": sig_items[0]["id"],
            "mode": "placeholder",
            "placeholder": "[[signature]]",
            "width": 144,
            "label": "replace placeholder with signature image",
        }
    else:
        placement = {
            "document": documents[0]["id"],
            "signature": sig_items[0]["id"],
            "page": 1,
            "anchor": "bottom-left",
            "x": 72,
            "y": 72,
            "width": 144,
            "label": "replace with interviewed placement",
        }

    config = {
        "version": 1,
        "customization_instructions": args.instruction or [],
        "documents": documents,
        "signatures": sig_items,
        "placements": [placement],
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(config, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {output_path}")
    if args.bundle_signatures:
        bundle_args = argparse.Namespace(
            config=str(output_path),
            asset_dir=None,
            instruction=[],
            signature_instruction=[],
        )
        command_bundle_signatures(bundle_args)
    return 0


def parse_signature_instructions(values: list[str]) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}
    for value in values:
        if "=" not in value:
            raise ConfigError(
                "signature instructions must use <signature-id>=<instruction>"
            )
        sig_id, instruction = value.split("=", 1)
        sig_id = sig_id.strip()
        instruction = instruction.strip()
        if not sig_id or not instruction:
            raise ConfigError(
                "signature instructions must use <signature-id>=<instruction>"
            )
        result.setdefault(sig_id, []).append(instruction)
    return result


def command_bundle_signatures(args: argparse.Namespace) -> int:
    config_path = Path(args.config).expanduser().resolve()
    data = read_json(config_path)
    signatures_raw = data.get("signatures")
    if not isinstance(signatures_raw, list) or not signatures_raw:
        raise ConfigError("signatures must be a non-empty array")

    if args.asset_dir:
        asset_dir = resolve_path(config_path, args.asset_dir)
    else:
        asset_dir = config_path.with_name(f"{config_path.stem}-assets") / "signatures"
    asset_dir.mkdir(parents=True, exist_ok=True)

    if args.instruction:
        existing = data.setdefault("customization_instructions", [])
        validate_string_list(existing, "customization_instructions")
        existing.extend(args.instruction)

    per_signature = parse_signature_instructions(args.signature_instruction or [])
    seen: set[str] = set()
    for index, item in enumerate(signatures_raw):
        context = f"signatures[{index}]"
        if not isinstance(item, dict):
            raise ConfigError(f"{context}: must be an object")
        sig_id = require_string(item, "id", context)
        seen.add(sig_id)
        source = resolve_path(config_path, require_string(item, "path", context))
        if not source.exists():
            raise ConfigError(f"{context}: signature image not found: {source}")
        target = asset_dir / f"{safe_asset_name(sig_id)}{source.suffix.lower()}"
        if source.resolve() != target.resolve():
            shutil.copy2(source, target)
        item["path"] = portable_path(config_path, target)
        if sig_id in per_signature:
            existing = item.setdefault("customization_instructions", [])
            validate_string_list(existing, f"{context}.customization_instructions")
            existing.extend(per_signature[sig_id])

    unknown = sorted(set(per_signature) - seen)
    if unknown:
        raise ConfigError(f"unknown signature id(s): {', '.join(unknown)}")

    config_path.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")
    print(f"Bundled {len(signatures_raw)} signature image(s) into {asset_dir}")
    print(f"Updated {config_path}")
    return 0


def command_validate(args: argparse.Namespace) -> int:
    config = validate_with_pdf_metadata(Path(args.config))
    docx_count = sum(1 for document in config.documents.values() if document.kind == "docx")
    print(
        f"Config valid: {len(config.documents)} document(s), "
        f"{len(config.signatures)} signature image(s), "
        f"{len(config.placements)} placement(s)"
    )
    if docx_count:
        print(f"Includes {docx_count} Word document(s); run apply for anchor checks.")
    return 0


def iter_docx_paragraphs(container: Any):
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_docx_paragraphs(cell)


def clear_paragraph(paragraph: Any) -> None:
    if not paragraph.runs:
        paragraph.add_run("")
    first = True
    for run in paragraph.runs:
        if first:
            run.text = ""
            first = False
        else:
            run.text = ""


def insert_paragraph_after(paragraph: Any) -> Any:
    from docx.text.paragraph import Paragraph  # type: ignore
    from docx.oxml import OxmlElement  # type: ignore

    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def add_docx_picture(paragraph: Any, image_path: Path, width_points: float) -> None:
    _, Inches = load_docx()
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_points / 72.0))


def apply_docx_placement(document: Any, placement: PlacementSpec, signature: SignatureSpec) -> None:
    mode = placement.mode or "append"
    if mode == "append":
        paragraph = document.add_paragraph()
        add_docx_picture(paragraph, signature.path, placement.width)
        return

    if mode == "placeholder":
        placeholder = placement.placeholder or ""
        for paragraph in iter_docx_paragraphs(document):
            text = paragraph.text
            if placeholder in text:
                before, after = text.split(placeholder, 1)
                clear_paragraph(paragraph)
                paragraph.runs[0].text = before
                add_docx_picture(paragraph, signature.path, placement.width)
                if after:
                    paragraph.add_run(after)
                return
        raise ConfigError(f"placeholder not found in DOCX: {placeholder!r}")

    if mode == "after-paragraph":
        target = placement.paragraph or ""
        for paragraph in iter_docx_paragraphs(document):
            if target in paragraph.text:
                new_paragraph = insert_paragraph_after(paragraph)
                add_docx_picture(new_paragraph, signature.path, placement.width)
                return
        raise ConfigError(f"paragraph text not found in DOCX: {target!r}")

    raise ConfigError(f"unsupported DOCX placement mode: {mode}")


def command_apply(args: argparse.Namespace) -> int:
    config_path = Path(args.config).expanduser().resolve()
    config = validate_with_pdf_metadata(config_path)
    fitz = load_fitz() if any(document.kind == "pdf" for document in config.documents.values()) else None
    placements_by_doc: dict[str, list[PlacementSpec]] = {}
    for placement in config.placements:
        placements_by_doc.setdefault(placement.document, []).append(placement)

    for doc_id, placements in placements_by_doc.items():
        document = config.documents[doc_id]
        output = document.output
        output.parent.mkdir(parents=True, exist_ok=True)
        if document.kind == "pdf":
            if fitz is None:
                raise ConfigError("PyMuPDF is required for PDF documents")
            pdf = fitz.open(str(document.path))
            try:
                for placement in placements:
                    if placement.page is None:
                        raise ConfigError("PDF placements require a page")
                    page = pdf[placement.page - 1]
                    signature = config.signatures[placement.signature]
                    image_w, image_h = image_size(fitz, signature.path)
                    rect = placement_rect(
                        fitz, page.rect.height, image_w, image_h, placement
                    )
                    page.insert_image(
                        rect,
                        filename=str(signature.path),
                        overlay=True,
                        rotate=placement.rotation,
                    )
                pdf.save(str(output), garbage=4, deflate=True)
                print(f"Wrote {output}")
            finally:
                pdf.close()
        elif document.kind == "docx":
            docx, _ = load_docx()
            word = docx.Document(str(document.path))
            for placement in placements:
                signature = config.signatures[placement.signature]
                apply_docx_placement(word, placement, signature)
            word.save(str(output))
            print(f"Wrote {output}")
        else:
            raise ConfigError(f"unsupported document type: {document.kind}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Create, validate, and apply visible PDF signature placements."
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    init_config = subparsers.add_parser("init-config")
    init_config.add_argument("--pdf", action="append", required=True)
    init_config.add_argument("--signature", action="append", required=True)
    init_config.add_argument("--output", required=True)
    init_config.add_argument("--instruction", action="append")
    init_config.add_argument("--bundle-signatures", action="store_true")
    init_config.set_defaults(func=command_init_config)

    bundle = subparsers.add_parser("bundle-signatures")
    bundle.add_argument("--config", required=True)
    bundle.add_argument("--asset-dir")
    bundle.add_argument("--instruction", action="append")
    bundle.add_argument("--signature-instruction", action="append")
    bundle.set_defaults(func=command_bundle_signatures)

    validate = subparsers.add_parser("validate")
    validate.add_argument("--config", required=True)
    validate.set_defaults(func=command_validate)

    apply = subparsers.add_parser("apply")
    apply.add_argument("--config", required=True)
    apply.set_defaults(func=command_apply)

    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        return args.func(args)
    except ConfigError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
